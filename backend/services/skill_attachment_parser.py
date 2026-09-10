"""Basic, non-executing parsers for personal Skill attachments."""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

import pymupdf
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md", ".py", ".js", ".sql", ".sh"}
MAX_PDF_PAGES = 100
MAX_XLSX_SHEETS = 10
MAX_XLSX_ROWS = 1000
MAX_XLSX_COLUMNS = 50
MAX_CSV_ROWS = 5000
MAX_CSV_COLUMNS = 100
TARGET_CHUNK_CHARS = 1600


class AttachmentParseError(ValueError):
    """The attachment cannot be converted into useful text."""


@dataclass(frozen=True)
class ParsedAttachment:
    chunks: list[dict[str, str | int]]
    page_count: int | None = None


def _clean_text(value: str) -> str:
    lines = [re.sub(r"[ \t\f\v]+", " ", line).strip() for line in value.replace("\x00", "").splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _split_text(text: str, location: str, chunks: list[dict[str, str | int]]) -> None:
    remaining = _clean_text(text)
    while remaining:
        if len(remaining) <= TARGET_CHUNK_CHARS:
            piece, remaining = remaining, ""
        else:
            split_at = remaining.rfind("\n", 0, TARGET_CHUNK_CHARS)
            if split_at < TARGET_CHUNK_CHARS // 2:
                split_at = remaining.rfind("。", 0, TARGET_CHUNK_CHARS) + 1
            if split_at < TARGET_CHUNK_CHARS // 2:
                split_at = TARGET_CHUNK_CHARS
            piece, remaining = remaining[:split_at], remaining[split_at:]
        cleaned = piece.strip()
        if cleaned:
            chunks.append({"chunk_index": len(chunks), "source_location": location, "content": cleaned})


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _parse_plain_text(path: Path) -> ParsedAttachment:
    lines = _decode_text(path.read_bytes()).splitlines()
    chunks: list[dict[str, str | int]] = []
    buffer: list[str] = []
    start_line = 1
    current_chars = 0
    for line_no, line in enumerate(lines, 1):
        if buffer and current_chars + len(line) + 1 > TARGET_CHUNK_CHARS:
            _split_text("\n".join(buffer), f"第 {start_line}-{line_no - 1} 行", chunks)
            buffer, start_line, current_chars = [], line_no, 0
        buffer.append(line)
        current_chars += len(line) + 1
    if buffer:
        _split_text("\n".join(buffer), f"第 {start_line}-{len(lines)} 行", chunks)
    return ParsedAttachment(chunks)


def _parse_pdf(path: Path) -> ParsedAttachment:
    chunks: list[dict[str, str | int]] = []
    with pymupdf.open(path) as document:
        page_count = len(document)
        for page_index in range(min(page_count, MAX_PDF_PAGES)):
            _split_text(document[page_index].get_text("text"), f"第 {page_index + 1} 页", chunks)
    return ParsedAttachment(chunks, page_count=page_count)


def _parse_docx(path: Path) -> ParsedAttachment:
    chunks: list[dict[str, str | int]] = []
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraph_no = 0
    for paragraph in root.iter(namespace + "p"):
        text = "".join(node.text or "" for node in paragraph.iter(namespace + "t"))
        if _clean_text(text):
            paragraph_no += 1
            _split_text(text, f"第 {paragraph_no} 段", chunks)
    return ParsedAttachment(chunks)


def _row_text(values: tuple[object, ...], max_columns: int) -> str:
    return "\t".join("" if value is None else str(value) for value in values[:max_columns]).rstrip()


def _parse_xlsx(path: Path) -> ParsedAttachment:
    chunks: list[dict[str, str | int]] = []
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for worksheet in workbook.worksheets[:MAX_XLSX_SHEETS]:
            rows: list[str] = []
            start_row = 1
            for row_no, values in enumerate(worksheet.iter_rows(values_only=True), 1):
                if row_no > MAX_XLSX_ROWS:
                    break
                text = _row_text(values, MAX_XLSX_COLUMNS)
                if not rows:
                    start_row = row_no
                rows.append(text)
                if len(rows) >= 25:
                    _split_text("\n".join(rows), f"工作表 {worksheet.title}，第 {start_row}-{row_no} 行", chunks)
                    rows = []
            if rows:
                end_row = min(worksheet.max_row or start_row, MAX_XLSX_ROWS)
                _split_text("\n".join(rows), f"工作表 {worksheet.title}，第 {start_row}-{end_row} 行", chunks)
    finally:
        workbook.close()
    return ParsedAttachment(chunks)


def _parse_csv(path: Path) -> ParsedAttachment:
    text = _decode_text(path.read_bytes())
    chunks: list[dict[str, str | int]] = []
    rows: list[str] = []
    start_row = 1
    for row_no, row in enumerate(csv.reader(io.StringIO(text)), 1):
        if row_no > MAX_CSV_ROWS:
            break
        if not rows:
            start_row = row_no
        rows.append("\t".join(row[:MAX_CSV_COLUMNS]))
        if len(rows) >= 50:
            _split_text("\n".join(rows), f"第 {start_row}-{row_no} 行", chunks)
            rows = []
    if rows:
        _split_text("\n".join(rows), f"第 {start_row}-{min(row_no, MAX_CSV_ROWS)} 行", chunks)
    return ParsedAttachment(chunks)


def parse_attachment(path: Path, original_filename: str) -> ParsedAttachment:
    """Parse one supported file without executing embedded code or macros."""
    suffix = Path(original_filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise AttachmentParseError("不支持解析该附件格式")
    try:
        if suffix == ".pdf":
            result = _parse_pdf(path)
        elif suffix == ".docx":
            result = _parse_docx(path)
        elif suffix == ".xlsx":
            result = _parse_xlsx(path)
        elif suffix == ".csv":
            result = _parse_csv(path)
        else:
            result = _parse_plain_text(path)
    except AttachmentParseError:
        raise
    except Exception as exc:
        raise AttachmentParseError(f"附件解析失败：{type(exc).__name__}") from exc
    if not result.chunks:
        message = "PDF 未提取到文字，可能是扫描版文件" if suffix == ".pdf" else "附件未提取到可用文字"
        raise AttachmentParseError(message)
    return result
